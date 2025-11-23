#!/usr/bin/env python3
"""Convert PROJECT_SUMMARY.md to Word docx format"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

def parse_and_add_content(doc, content):
    """Parse markdown and add to Word document"""
    lines = content.split('\n')

    for line in lines:
        # Skip empty lines
        if not line.strip():
            continue

        # H1 headers
        if line.startswith('# '):
            p = doc.add_heading(line[2:].strip(), level=1)

        # H2 headers
        elif line.startswith('## '):
            p = doc.add_heading(line[3:].strip(), level=2)

        # H3 headers
        elif line.startswith('### '):
            p = doc.add_heading(line[4:].strip(), level=3)

        # H4 headers
        elif line.startswith('#### '):
            p = doc.add_heading(line[5:].strip(), level=4)

        # Horizontal rule
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.add_run('_' * 80)

        # Numbered lists
        elif re.match(r'^\d+\.', line.strip()):
            # Remove markdown bold markers and parse
            clean_line = line.strip()
            p = doc.add_paragraph(style='List Number')

            # Parse bold text
            parts = re.split(r'(\*\*.*?\*\*)', clean_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    # Remove leading number and dot
                    text = re.sub(r'^\d+\.\s*', '', part)
                    if text:
                        p.add_run(text)

        # Bullet points
        elif line.strip().startswith('- ') or line.strip().startswith('* '):
            clean_line = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')

            # Parse bold text
            parts = re.split(r'(\*\*.*?\*\*)', clean_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

        # Bold text (starts with **)
        elif line.strip().startswith('**'):
            p = doc.add_paragraph()
            # Parse bold and regular text
            parts = re.split(r'(\*\*.*?\*\*)', line.strip())
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)

        # Regular paragraphs
        else:
            clean_line = line.strip()
            if clean_line:
                p = doc.add_paragraph()
                # Parse bold text
                parts = re.split(r'(\*\*.*?\*\*)', clean_line)
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        p.add_run(part)

def main():
    # Read the markdown file
    with open('PROJECT_SUMMARY.md', 'r', encoding='utf-8') as f:
        content = f.read()

    # Create Word document
    doc = Document()

    # Set document title
    title = doc.add_heading('SonicSoothe: AI-Powered Music Therapy Platform', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Parse and add content
    parse_and_add_content(doc, content)

    # Save document
    doc.save('PROJECT_SUMMARY.docx')
    print("Successfully converted PROJECT_SUMMARY.md to PROJECT_SUMMARY.docx")

if __name__ == "__main__":
    main()
